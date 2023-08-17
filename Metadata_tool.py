#import required modules
import re
import requests
from collections import deque
from urllib.parse import urlsplit
from bs4 import BeautifulSoup
import pandas as pd
from PyPDF2 import PdfFileReader
import docx
from openpyxl import load_workbook
import pptx
from googlesearch import search

#send request to target site
while True:
    original_url = input("Enter the website url(in 'http://www.example.domain' format only):  ")

    try:
        r = requests.get(original_url)
        status = int(r.status_code)
        if status not in range(200, 300):
            print ("Site unavailable or does not exist. Please enter URL again.\n")
        else:
            break
    except (requests.exceptions.InvalidSchema, requests.exceptions.InvalidURL, requests.exceptions.ConnectionError, requests.exceptions.MissingSchema):
        print("Error. Try again.\n")
        continue
    
unscraped = deque([original_url])
scraped = []
docs = []
new_docs=[]
new_links = []


split = urlsplit(original_url)
loc = ("{0.netloc}".format(split)).split('.')[1]
file = loc + ".xlsx"
base_url = "{0.scheme}://{0.netloc}".format(split)

if '/' in split.path:
    path = original_url[:original_url.rfind('/')+1]
else:
    path = original_url

     
main_response = requests.get(base_url)

soup = BeautifulSoup(main_response.text, 'html.parser')

#find links matching required document format
for anchor in soup.find_all("a"):
    if "href" in anchor.attrs:
        link = anchor.attrs["href"]

    else:
        link = ''
    if link.startswith('/'):
        link = base_url +link
        
    elif not link.startswith('http'):
        link = path + link
        
    if not link in unscraped and not link in scraped:
        unscraped.append(link)
    if link.endswith('.pdf') or link.endswith('.doc') or link.endswith('.docx') or link.endswith('.xls') or link.endswith('.xlsx') or link.endswith('.ppt') or link.endswith('.pptx'):
        new_links.append(link)

#print out links of documents found       
print(new_links)

#pdf metadata
def pdf_data(path):
    with open(path, 'rb') as f:
        pdf = PdfFileReader(f, strict=False)
        information = pdf.getDocumentInfo()
        number_of_pages = pdf.getNumPages()

    txt = f"""
    Information about {path}: 

    Author: {information.author}
    Creator: {information.creator}
    Producer: {information.producer}
    Subject: {information.subject}
    Title: {information.title}
    Number of pages: {number_of_pages}
    """

    return(txt)

#doc metadata
def doc_data(doc):
    prop = doc.core_properties
    
    txt = f"""
    Information about {doc}: 

    Author: {prop.author}
    Category: {prop.category}
    Title: {prop.title}
    Subject: {prop.subject}
    Version: {prop.version}
    Identifier: {prop.identifier}
    Keywords: {prop.keywords}
    Content status: {prop.content_status}
    Created: str({prop.created})
    Modified: str({prop.modified})
    Language: {prop.language}
    """
    return(txt)

#excel metadata
def excel_data(xl):
    wb = load_workbook(xl)
    prop = wb.properties
    txt = f"""
    Information about {xl}: 

    Author: {prop.creator}
    Category: {prop.category}
    Title: {prop.title}
    Subject:{prop.subject}
    Description: {prop.description}
    Version: {prop.version}
    Identifier: {prop.identifier}
    Keywords: {prop.keywords}
    Created: {prop.created}
    Modified: {prop.modified}
    Last modified by: {prop.lastModifiedBy}
    Language: {prop.language}
    
    """
    return (txt)
#ppt metadata
def ppt_data(file):
    ppt = pptx.Presentation(file)
    
    prop = ppt.core_properties

    txt = f"""
    Information about {file}: 

    Author: {prop.author}
    Category: {prop.category}
    Title: {prop.title}
    Subject: {prop.subject}
    Version: {prop.version}
    Identifier: {prop.identifier}
    Keywords: {prop.keywords}
    Content status: {prop.content_status}
    Created: str({prop.created})
    Modified: str({prop.modified})
    Language: {prop.language}
    
    """
    return(txt)

#download files
for i in new_links:
    print("Downloading 1 file...")
    req = requests.get(i)
    filename = req.url[i.rfind('/')+1:]
    new_docs.append(filename)
    if i.startswith('https://') or i.startswith('http://'):
        req = requests.get(i)
        open(filename,'wb').write(req.content)
print("Files downloaded")


#extract metadata for files to a .txt file 
for j in new_docs:
    if j.endswith('.pdf'):
        f = open("demopdf.txt","a")
        f.write(str(j) + ': \n' + pdf_data(j) + '\n')
        f.close()
    elif j.endswith('.doc') or j.endswith('.docx'):
        g = open("demoword.txt","a")
        f.write(j + ': \n' + doc_data(j) + '\n')
        f.close()
    elif j.endswith('.xls') or j.endswith('.xlsx'):
        f = open("demoxls.txt","a")
        f.write(j + ': \n' + excel_data(j) + '\n')
        f.close()
    elif j.endswith('.ppt') or j.endswith('.pptx'):
        f = open("demoppt.txt",'a')
        f.write(str(j) + ': \n' + ppt_data(j) + '\n')
        f.close

print("Metadata extraction complete")